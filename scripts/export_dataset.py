"""Đóng băng bộ dữ liệu: Excel -> data/v1.0/{dev,test,unsure}.jsonl + taxonomy.json + MANIFEST.

Chạy: python scripts/export_dataset.py <đường-dẫn-xlsx> <thư-mục-ra>
"""
import json, re, sys, hashlib, os
from collections import Counter
from openpyxl import load_workbook

XLSX, OUT = sys.argv[1], sys.argv[2]
VERSION = "1.0"
os.makedirs(OUT, exist_ok=True)

wb = load_workbook(XLSX, data_only=True)
d = wb["Dữ liệu"]
h = [c.value for c in d[1]]
SIG = [k for k in h if k and re.match(r"S\d\d_", k)]

rows = []
for r in d.iter_rows(min_row=2, values_only=True):
    x = dict(zip(h, r))
    if not x["id"]:
        continue
    assert x["check"] == "OK", (x["id"], x["check"])
    rec = {
        "id": x["id"],
        "text": x["text"],
        "source_type": x["source_type"],
        "language": x["language"],
        "label": x["label"],
        "signals": [s[:3] for s in SIG if x[s] == 1],
        "hard_negative": (x["hard_negative"] == "yes") if x["label"] == "legit" else None,
        "split": x["split"],
        "source": x["source"],
        "label_notes": x["notes"],
    }
    rows.append(rec)

parts = {"dev": [r for r in rows if r["split"] == "dev"],
         "test": [r for r in rows if r["split"] == "test"],
         "unsure": [r for r in rows if r["label"] == "unsure"]}
assert len(parts["dev"]) + len(parts["test"]) + len(parts["unsure"]) == len(rows)
assert not any(r["split"] for r in parts["unsure"])

for name, rs in parts.items():
    with open(f"{OUT}/{name}.jsonl", "w", encoding="utf-8") as f:
        for r in rs:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")

t = wb["Taxonomy"]
tax = {"version": VERSION, "signals": []}
for r in t.iter_rows(min_row=2, values_only=True):
    if r[0] and re.match(r"S\d\d$", str(r[0])):
        tax["signals"].append({"id": r[0], "column": r[1], "name": r[2], "definition": r[3],
                               "cues": r[4], "example": r[5], "exclusions": r[6]})
assert len(tax["signals"]) == 14
with open(f"{OUT}/taxonomy.json", "w", encoding="utf-8") as f:
    json.dump(tax, f, ensure_ascii=False, indent=2)

with open(f"{OUT}/MANIFEST.sha256", "w") as f:
    for name in ["dev.jsonl", "test.jsonl", "unsure.jsonl", "taxonomy.json"]:
        f.write(hashlib.sha256(open(f"{OUT}/{name}", "rb").read()).hexdigest() + f"  {name}\n")

# số liệu cho DATASET.md
stats = {k: len(v) for k, v in parts.items()}
by = Counter((r["split"], r["label"], r["source_type"]) for r in rows if r["split"])
sig = Counter((r["split"], r["label"], s) for r in rows if r["split"] for s in r["signals"])
hn = Counter(r["split"] for r in rows if r["hard_negative"])
json.dump({"stats": stats, "by": {"|".join(k): v for k, v in by.items()},
           "sig": {"|".join(k): v for k, v in sig.items()}, "hard_negative": hn},
          open(f"{OUT}/../stats_v{VERSION}.json", "w"), ensure_ascii=False, indent=1)
print(stats, dict(hn))
