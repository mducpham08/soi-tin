"""Chia dev/test (~40/60) cho scam-detector-dataset-v3.xlsx.

Ràng buộc:
1. Các tin cùng nhóm (cùng hội thoại / chiến dịch / khuôn mẫu / gần trùng) đi nguyên khối vào một tập.
2. Dòng có ghi chú "CHỈ đưa vào tập TEST" (gần trùng test.csv của nguồn) bắt buộc vào test, kéo theo cả nhóm.
3. Phân tầng: tỉ lệ dev ~40% đồng đều theo label, source_type, hard_negative, ngôn ngữ và từng dấu hiệu S01–S14.
4. Dòng unsure không chia (để trống).
Tái lập được: seed cố định.
"""
import re, random, sys, math
from collections import defaultdict, Counter
from openpyxl import load_workbook

FN = sys.argv[1] if len(sys.argv) > 1 else "scam-detector-dataset-v3.xlsx"
DEV_FRAC, SEED, JAC = 0.40, 42, 0.5

wb = load_workbook(FN); d = wb["Dữ liệu"]
h = [c.value for c in d[1]]; col = {k: i + 1 for i, k in enumerate(h)}
SIG = [k for k in h if k and re.match(r"S\d\d_", k)]
rows = []
for r in range(2, d.max_row + 1):
    rid = d.cell(r, 1).value
    if not rid:
        continue
    x = {k: d.cell(r, col[k]).value for k in h if k}
    x["_r"] = r
    rows.append(x)
use = [x for x in rows if x["label"] in ("scam", "legit")]
byid = {x["id"]: x for x in use}

# ---------- nhóm (union-find) ----------
par = {x["id"]: x["id"] for x in use}
def find(a):
    while par[a] != a:
        par[a] = par[par[a]]; a = par[a]
    return a
def union(a, b):
    if a in par and b in par:
        par[find(a)] = find(b)

MANUAL = [
    ["DOC003", "DOC004", "DOC005", "DOC006", "DOC007"],          # mẫu thư 'Quyết định số 04'
    ["DOC009", "DOC010"],                                        # mẫu thư trao đổi SV tại Úc
    ["SMS094", "SMS097"],                                        # [B00king_CTV]
    ["CHAT001", "DOC017"],                                       # cùng bài Thanh Niên (Circle K)
    ["JOB025", "CHAT002", "CHAT003"],                            # chiến dịch Mega Market
    ["EMAIL005", "DOC018", "DOC019", "DOC020", "JOB030"],        # Vinamilk + khuôn 'Quyền lợi' Vinbus
    ["EMAIL003", "EMAIL004"],                                    # Tâm Anh
]
for g in MANUAL:
    for a in g[1:]:
        union(g[0], a)

def toks(s):
    return set(re.findall(r"\w+", str(s).lower())) - {"money", "number", "time", "date", "ten", "sdt", "stk", "otp", "point"}
tk = {x["id"]: toks(x["text"]) for x in use}
tag = {x["id"]: (re.match(r"^\[([^\]]+)\]", str(x["text"])) or [None, None])[1] for x in use}
ids = list(byid)
for i, a in enumerate(ids):
    for b in ids[i + 1:]:
        A, B = tk[a], tk[b]
        if A and B and len(A & B) / len(A | B) >= JAC:
            union(a, b)
        if tag[a] and tag[a] == tag[b] and "_" in tag[a]:
            union(a, b)
# tin tuyển dụng legit cùng công ty (SGU) → cùng nhóm
for kw in ["WESET", "Annam Gourmet"]:
    g = [x["id"] for x in use if x["source_type"] == "job_post" and kw in str(x["text"])]
    for a in g[1:]:
        union(g[0], a)

groups = defaultdict(list)
for x in use:
    groups[find(x["id"])].append(x["id"])
forced_test = {find(x["id"]) for x in use if "CHỈ đưa vào tập TEST" in str(x["notes"] or "")}

# ---------- đặc trưng để phân tầng ----------
def feats(x):
    f = [f"label={x['label']}", f"src={x['source_type']}|{x['label']}", f"lang={x['language']}"]
    if x["hard_negative"] == "yes":
        f.append("hard_negative")
    for s in SIG:
        if x[s] == 1:
            f.append(f"{s[:3]}|{x['label']}")
    return f
F = {x["id"]: feats(x) for x in use}
tot = Counter(f for x in use for f in F[x["id"]])
N = len(use)

def cost(devc, ndev):
    c = ((ndev - DEV_FRAC * N) / math.sqrt(N)) ** 2 * 4
    for f, t in tot.items():
        c += ((devc[f] - DEV_FRAC * t) / math.sqrt(t)) ** 2
    return c

def gfeat(g):
    return Counter(f for i in groups[g] for f in F[i])

rng = random.Random(SEED)
order = [g for g in groups if g not in forced_test]
rng.shuffle(order)
order.sort(key=lambda g: -len(groups[g]))      # nhóm lớn xếp trước
assign = {g: "test" for g in forced_test}
devc, ndev = Counter(), 0
for g in order:
    gf, n = gfeat(g), len(groups[g])
    c_dev = cost(devc + gf, ndev + n)
    c_test = cost(devc, ndev)
    if c_dev < c_test:
        assign[g] = "dev"; devc += gf; ndev += n
    else:
        assign[g] = "test"

# tìm kiếm cục bộ: đổi tập từng nhóm nếu giảm cost
improved = True
while improved:
    improved = False
    for g in order:
        gf, n = gfeat(g), len(groups[g])
        if assign[g] == "dev":
            nc = cost(devc - gf, ndev - n)
            if nc < cost(devc, ndev) - 1e-9:
                assign[g] = "test"; devc -= gf; ndev -= n; improved = True
        else:
            nc = cost(devc + gf, ndev + n)
            if nc < cost(devc, ndev) - 1e-9:
                assign[g] = "dev"; devc += gf; ndev += n; improved = True

split = {i: assign[find(i)] for i in byid}

# ---------- ghi ----------
for x in rows:
    d.cell(x["_r"], col["split"]).value = split.get(x["id"])   # unsure -> None
wb.save(FN)

# ---------- báo cáo ----------
multi = sorted([sorted(v) for v in groups.values() if len(v) > 1], key=len, reverse=True)
print("Nhóm có >1 dòng:", len(multi))
for v in multi:
    print(f"  [{split[v[0]]}] " + ", ".join(v))
dv = Counter(split.values())
print("dev/test:", dv["dev"], dv["test"])
print(f"{'đặc trưng':28s} {'tổng':>5s} {'dev':>4s} {'test':>4s} {'%dev':>5s}")
for f, t in sorted(tot.items()):
    k = sum(1 for i in byid if f in F[i] and split[i] == "dev")
    print(f"{f:28s} {t:5d} {k:4d} {t-k:4d} {k/t:5.0%}")
print("forced test:", sorted(i for g in forced_test for i in groups[g]))
